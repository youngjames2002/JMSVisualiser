import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from io import BytesIO, StringIO
import msal
import requests
import psycopg2

# read scheduling db
def get_connection():
    return psycopg2.connect(st.secrets["SCHEDULEDB"]["DATABASE_URL"])

def load_scheduling_data():
    conn = get_connection()
    scheduling_df = pd.read_sql_query("SELECT * FROM schedule;", conn)
    return scheduling_df

# read bundles data
def load_bundles_table():
    conn = get_connection()
    df = pd.read_sql_query("SELECT * FROM bundles;", conn)
    return df

# join scheduling table to bundles table
def join_bundles_to_schedule(scheduling_df, bundles_df):
    flat_bundle_names = set(bundles_df[bundles_df["type"] == "FLAT"]["bundlename"])
    tube_bundle_names = set(bundles_df[bundles_df["type"] == "TUBE"]["bundlename"])
    so_flat = (scheduling_df[scheduling_df["bundle"].isin(flat_bundle_names)]
               .groupby("sonumber")["bundle"].first().rename("flat_bundle"))
    so_tube = (scheduling_df[scheduling_df["bundle"].isin(tube_bundle_names)]
               .groupby("sonumber")["bundle"].first().rename("tube_bundle"))
    scheduling_df = scheduling_df.merge(so_flat, on="sonumber", how="left")
    scheduling_df = scheduling_df.merge(so_tube, on="sonumber", how="left")
    df = scheduling_df.merge(bundles_df, left_on="bundle", right_on="bundlename", how="left").drop(columns="bundlename")
    return df

# convert to format visualiser expects?

# read bmena finishing schedule sheet
@st.cache_data(show_spinner=True)
def download_excel_from_sharepoint(site_name: str, file_path:str) -> BytesIO:
    # download from sharepoint and return bytesIO object

    TENANT_ID = st.secrets["sharepoint"]["TENANT_ID"]
    CLIENT_ID = st.secrets["sharepoint"]["CLIENT_ID"]
    CLIENT_SECRET = st.secrets["sharepoint"]["CLIENT_SECRET"]
    SHAREPOINT_SITE = st.secrets["sharepoint"]["SHAREPOINT_SITE"]

    AUTHORITY = f"https://login.microsoftonline.com/{TENANT_ID}"
    SCOPE = ["https://graph.microsoft.com/.default"]

    app = msal.ConfidentialClientApplication(
        client_id=CLIENT_ID,
        authority=AUTHORITY,
        client_credential=CLIENT_SECRET
    )  

    token = app.acquire_token_for_client(scopes=SCOPE)
    if "access_token" not in token:
        st.error("Authentication failed")
        return None

    headers = {"Authorization": f"Bearer {token['access_token']}"} 

    # Get SharePoint site ID
    site_url = f"https://graph.microsoft.com/v1.0/sites/{SHAREPOINT_SITE}:/sites/{site_name}:/"
    site_response = requests.get(site_url, headers=headers)

    if site_response.status_code != 200:
        st.error("Site lookup failed")
        return None

    site_id = site_response.json()["id"]

    # Download the file
    file_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_path}:/content"
    file_response = requests.get(file_url, headers=headers)
    file_response.raise_for_status()

    return BytesIO(file_response.content)

@st.cache_data(show_spinner=True)
def load_data_Bmena_sp():
    bytes_io = download_excel_from_sharepoint(
        site_name="JMSEngineeringTeam",
        file_path="JMS Engineering Team SharePoint/Paint Schedule/Bmena Finishing Schedule.xlsm"
    )
    if bytes_io is None:
        return pd.DataFrame()  # return empty DataFrame if download failed

    wb = load_workbook(filename=bytes_io, data_only=True)
    sheet = wb["Schedule"]
    lookup_table = sheet.tables["Table1"]
    data = sheet[lookup_table.ref]

    # Convert table to DataFrame
    rows_list = [[cell.value for cell in row] for row in data]
    df = pd.DataFrame(rows_list[1:], columns=rows_list[0])
    return df